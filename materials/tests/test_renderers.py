import csv
from datetime import datetime
from decimal import Decimal
from io import BytesIO, StringIO
from zoneinfo import ZoneInfo

from django.test import SimpleTestCase, TestCase, override_settings
from django.utils import timezone
from openpyxl import load_workbook

from utils.properties.models import Unit

from ..models import (
    ComponentMeasurement,
    Material,
    MaterialComponent,
    MaterialComponentGroup,
    MaterialProperty,
    MaterialPropertyValue,
    Sample,
)
from ..renderers import (
    SampleCSVRenderer,
    SampleMeasurementsXLSXRenderer,
    SampleXLSXRenderer,
)


class SamplePrecisionExportRendererTestCase(SimpleTestCase):
    def test_csv_and_xlsx_fixed_headers_export_stored_sampling_precision(self):
        data = [
            {
                "name": f"Precision {precision or 'legacy'}",
                "datetime": "2024-08-26T22:00:00Z",
                "datetime_precision": precision,
            }
            for precision in ("", "year", "date", "time")
        ]
        for renderer_class in (SampleCSVRenderer, SampleXLSXRenderer):
            with self.subTest(renderer=renderer_class):
                renderer = renderer_class()
                self.assertEqual(
                    renderer.labels.get("datetime_precision"), "Sampling date precision"
                )
                buffer = BytesIO()
                renderer.render(buffer, data)
                if renderer_class is SampleCSVRenderer:
                    rows = list(
                        csv.reader(StringIO(buffer.getvalue().decode("utf-8-sig")))
                    )
                else:
                    buffer.seek(0)
                    workbook = load_workbook(buffer)
                    rows = list(workbook.active.iter_rows(values_only=True))
                header = rows[0]
                precision_column = header.index("Sampling date precision")
                datetime_column = header.index("Sampling date/time")
                for row, expected in zip(rows[1:], data, strict=True):
                    self.assertEqual(
                        row[precision_column] or "", expected["datetime_precision"]
                    )
                    self.assertEqual(row[datetime_column], expected["datetime"])


class SampleMeasurementsXLSXRendererTestCase(TestCase):
    @override_settings(TIME_ZONE="Europe/Berlin", USE_TZ=True)
    def test_sampling_metadata_uses_precision_and_default_timezone(self):
        material = Material.objects.create(name="Dated renderer material")
        sample = Sample.objects.create(
            name="Dated renderer sample",
            material=material,
            analysis_date=datetime(2024, 9, 1, 13, 45, tzinfo=ZoneInfo("UTC")),
            analysis_laboratory="Renderer lab",
        )
        cases = (
            ("year", datetime(2024, 1, 1), "2024"),
            ("date", datetime(2024, 8, 27), "2024-08-27"),
            ("time", datetime(2024, 8, 27, 14, 30), "2024-08-27 14:30"),
            ("time", datetime(2024, 8, 27), "2024-08-27 00:00"),
            ("", datetime(2024, 8, 27), "2024-08-27"),
            ("", datetime(2024, 8, 27, 14, 30), "2024-08-27 14:30"),
            ("", None, ""),
        )
        for precision, value, expected in cases:
            with self.subTest(precision=precision, value=value):
                sample.datetime = (
                    value.replace(tzinfo=ZoneInfo("Europe/Berlin")).astimezone(
                        ZoneInfo("UTC")
                    )
                    if value
                    else None
                )
                sample.datetime_precision = precision
                sample.save(update_fields=["datetime", "datetime_precision"])
                sample.refresh_from_db()
                original = sample.datetime
                with timezone.override("America/Los_Angeles"):
                    buffer = SampleMeasurementsXLSXRenderer(
                        sample=sample,
                        measurements=sample.component_measurements.all(),
                    ).render()
                workbook = load_workbook(buffer)
                for sheet_name in ("Measurements", "Sample Properties"):
                    metadata = {
                        row[0]: row[1]
                        for row in workbook[sheet_name].iter_rows(
                            min_row=1, max_row=13, max_col=2, values_only=True
                        )
                    }
                    self.assertEqual(metadata["Sample date"] or "", expected)
                    self.assertEqual(metadata["Analysis date"], "2024-09-01")
                    self.assertEqual(metadata["Analysis laboratorium"], "Renderer lab")
                self.assertEqual(sample.datetime, original)

    def test_render_exports_raw_component_measurements(self):
        material = Material.objects.create(name="Renderer Material")
        sample = Sample.objects.create(name="Renderer Sample", material=material)
        group = MaterialComponentGroup.objects.create(name="Renderer Group")
        raw_component = MaterialComponent.objects.create(name="Raw Export Component")
        unit = Unit.objects.filter(name="%").first() or Unit.objects.create(name="%")

        ComponentMeasurement.objects.create(
            sample=sample,
            group=group,
            component=raw_component,
            unit=unit,
            average=Decimal("42"),
            standard_deviation=Decimal("1.5"),
        )

        buffer = SampleMeasurementsXLSXRenderer(
            sample=sample,
            measurements=sample.component_measurements.all(),
        ).render()
        workbook = load_workbook(buffer)
        measurements_sheet = workbook["Measurements"]
        exported_values = [
            cell.value for cell in measurements_sheet[16] if cell.value is not None
        ]

        self.assertIn("Raw Export Component", exported_values)
        self.assertIn(42, exported_values)

    def test_export_includes_qualifier_metadata_and_keeps_raw_text_literal(self):
        material = Material.objects.create(name="Qualifier Material")
        sample = Sample.objects.create(name="Qualifier Sample", material=material)
        group = MaterialComponentGroup.objects.create(name="Qualifier Group")
        unit = Unit.objects.filter(name="%").first() or Unit.objects.create(name="%")
        basis = MaterialComponent.objects.create(name="Qualifier Basis")

        def add_measurement(name, **kwargs):
            return ComponentMeasurement.objects.create(
                sample=sample,
                group=group,
                component=MaterialComponent.objects.create(name=name),
                basis_component=basis,
                unit=unit,
                **kwargs,
            )

        add_measurement("Exact Zero", average=Decimal("0"))
        add_measurement("Exact Formula", average=Decimal("42"), raw_value="=1+1")
        add_measurement(
            "Censored Less",
            average=Decimal("0.1"),
            raw_value="0.1",
            value_qualifier="less_than",
            detection_limit=Decimal("0"),
        )
        add_measurement(
            "Censored Greater",
            average=Decimal("1450"),
            raw_value=">1450",
            value_qualifier="greater_than",
            detection_limit=Decimal("6.57"),
            raw_detection_limit="=2+2",
        )
        add_measurement(
            "Below Limit",
            average=Decimal("0"),
            raw_value="(5,63)<LD=6,57",
            value_qualifier="below_detection_limit",
        )
        add_measurement(
            "Censored Formula",
            average=Decimal("1"),
            raw_value="=1+1",
            value_qualifier="less_than",
        )

        exact_property = MaterialProperty.objects.create(
            name="Exact Property", unit="%"
        )
        censored_property = MaterialProperty.objects.create(
            name="Censored Property", unit="%"
        )
        MaterialPropertyValue.objects.create(
            sample=sample,
            property=exact_property,
            unit=unit,
            average=Decimal("42"),
            raw_value="=1+1",
        )
        MaterialPropertyValue.objects.create(
            sample=sample,
            property=censored_property,
            unit=unit,
            average=Decimal("0.5"),
            raw_value="0.5",
            value_qualifier="less_than",
            detection_limit=Decimal("0"),
            raw_detection_limit="=3+3",
        )

        buffer = SampleMeasurementsXLSXRenderer(
            sample=sample,
            measurements=sample.component_measurements.all(),
        ).render()
        workbook = load_workbook(buffer, data_only=False)

        metadata_headers = [
            "Value qualifier",
            "Raw value",
            "Detection limit",
            "Raw detection limit",
        ]

        sheet = workbook["Measurements"]
        headers = [cell.value for cell in sheet[15]]
        self.assertEqual(headers[-4:], metadata_headers)
        value_col = headers.index("Value") + 1
        qualifier_col = headers.index("Value qualifier") + 1
        raw_col = headers.index("Raw value") + 1
        limit_col = headers.index("Detection limit") + 1
        raw_limit_col = headers.index("Raw detection limit") + 1
        rows = {
            sheet.cell(row=row, column=1).value: row
            for row in range(16, sheet.max_row + 1)
        }

        row = rows["Exact Zero"]
        value_cell = sheet.cell(row=row, column=value_col)
        self.assertEqual(value_cell.value, 0)
        self.assertNotIsInstance(value_cell.value, str)
        self.assertEqual(sheet.cell(row=row, column=qualifier_col).value, "exact")
        self.assertFalse(sheet.cell(row=row, column=raw_col).value)
        self.assertFalse(sheet.cell(row=row, column=limit_col).value)

        row = rows["Exact Formula"]
        self.assertEqual(sheet.cell(row=row, column=value_col).value, 42)
        raw_cell = sheet.cell(row=row, column=raw_col)
        self.assertEqual(raw_cell.value, "=1+1")
        self.assertEqual(raw_cell.data_type, "s")

        row = rows["Censored Less"]
        value_cell = sheet.cell(row=row, column=value_col)
        self.assertEqual(value_cell.value, "<0.1")
        self.assertEqual(value_cell.data_type, "s")
        self.assertEqual(sheet.cell(row=row, column=qualifier_col).value, "less_than")
        self.assertEqual(sheet.cell(row=row, column=raw_col).value, "0.1")
        self.assertEqual(sheet.cell(row=row, column=limit_col).value, 0)

        row = rows["Censored Greater"]
        self.assertEqual(sheet.cell(row=row, column=value_col).value, ">1450")
        self.assertEqual(sheet.cell(row=row, column=limit_col).value, 6.57)
        raw_limit_cell = sheet.cell(row=row, column=raw_limit_col)
        self.assertEqual(raw_limit_cell.value, "=2+2")
        self.assertEqual(raw_limit_cell.data_type, "s")

        row = rows["Below Limit"]
        self.assertEqual(
            sheet.cell(row=row, column=value_col).value,
            "Below detection limit ((5,63)<LD=6,57)",
        )
        self.assertEqual(
            sheet.cell(row=row, column=qualifier_col).value, "below_detection_limit"
        )
        self.assertFalse(sheet.cell(row=row, column=limit_col).value)
        self.assertFalse(sheet.cell(row=row, column=raw_limit_col).value)

        row = rows["Censored Formula"]
        value_cell = sheet.cell(row=row, column=value_col)
        self.assertEqual(value_cell.value, "<=1+1")
        self.assertEqual(value_cell.data_type, "s")

        sheet = workbook["Sample Properties"]
        headers = [cell.value for cell in sheet[15]]
        self.assertEqual(headers[-4:], metadata_headers)
        value_col = headers.index("Value") + 1
        qualifier_col = headers.index("Value qualifier") + 1
        raw_col = headers.index("Raw value") + 1
        limit_col = headers.index("Detection limit") + 1
        raw_limit_col = headers.index("Raw detection limit") + 1
        rows = {
            sheet.cell(row=row, column=1).value: row
            for row in range(16, sheet.max_row + 1)
        }

        row = rows["Exact Property"]
        self.assertEqual(sheet.cell(row=row, column=value_col).value, 42)
        raw_cell = sheet.cell(row=row, column=raw_col)
        self.assertEqual(raw_cell.value, "=1+1")
        self.assertEqual(raw_cell.data_type, "s")
        self.assertEqual(sheet.cell(row=row, column=qualifier_col).value, "exact")

        row = rows["Censored Property"]
        value_cell = sheet.cell(row=row, column=value_col)
        self.assertEqual(value_cell.value, "<0.5")
        self.assertEqual(value_cell.data_type, "s")
        self.assertEqual(sheet.cell(row=row, column=qualifier_col).value, "less_than")
        self.assertEqual(sheet.cell(row=row, column=limit_col).value, 0)
        raw_limit_cell = sheet.cell(row=row, column=raw_limit_col)
        self.assertEqual(raw_limit_cell.value, "=3+3")
        self.assertEqual(raw_limit_cell.data_type, "s")
